import hashlib
import re
import html
from io import BytesIO
from datetime import datetime

import fitz
import streamlit as st
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from document_reader import extract_text_from_pdf
from gemini_compare import compare_documents


# =========================================================
# PAGE CONFIGURATION
# =========================================================

MAX_PDF_SIZE_MB = 10
MAX_PDF_SIZE_BYTES = MAX_PDF_SIZE_MB * 1024 * 1024

st.set_page_config(
    page_title="DocuMind AI | Document Intelligence",
    page_icon="📘",
    layout="wide",
    initial_sidebar_state="expanded",
)


# =========================================================
# PROFESSIONAL STYLING
# =========================================================

st.markdown(
    """
    <style>
    .stApp {
        background-color: #0B1220;
        color: #F1F5F9;
    }

    [data-testid="stSidebar"] {
        background-color: #070D19;
        border-right: 1px solid #1E293B;
    }

    [data-testid="stSidebar"] * {
        color: #E2E8F0;
    }

    [data-testid="stHeader"],
    [data-testid="stToolbar"],
    [data-testid="stDecoration"] {
        background-color: #0B1220 !important;
    }

    [data-testid="stMainBlockContainer"] {
        max-width: 100% !important;
        padding-left: 3rem;
        padding-right: 3rem;
    }

    .hero {
        background: linear-gradient(135deg, #111D35, #1E40AF);
        padding: 32px;
        border: 1px solid #334F87;
        border-radius: 18px;
        margin-bottom: 25px;
    }

    .hero h1 {
        color: #FFFFFF;
        font-size: 36px;
        margin-bottom: 8px;
    }

    .hero p {
        color: #DBEAFE;
        font-size: 16px;
    }

    .section-heading {
        color: #93C5FD;
        font-size: 23px;
        font-weight: 700;
        margin: 20px 0 12px;
    }

    .info-card {
        background-color: #131E32;
        border: 1px solid #293A56;
        border-radius: 14px;
        padding: 18px;
        min-height: 110px;
        margin-bottom: 12px;
        overflow-wrap: anywhere;
    }

    .info-label {
        color: #94A3B8;
        font-size: 13px;
        margin-bottom: 8px;
    }

    .info-value {
        color: #F8FAFC;
        font-size: 20px;
        font-weight: 700;
        overflow-wrap: anywhere;
    }

    .footer {
        text-align: center;
        color: #94A3B8;
        font-size: 12px;
        padding: 25px 0 10px;
    }

    .finding {
        padding: 10px 14px;
        margin: 7px 0;
        border-radius: 7px;
        border-left: 4px solid;
        background-color: #131E32;
        line-height: 1.65;
        overflow-wrap: anywhere;
    }

    div.stButton > button,
    div[data-testid="stDownloadButton"] button {
        background-color: #2563EB;
        color: #FFFFFF;
        border: 1px solid #3B82F6;
        border-radius: 9px;
        font-weight: 600;
        min-height: 45px;
    }

    div.stButton > button:hover,
    div[data-testid="stDownloadButton"] button:hover {
        background-color: #1D4ED8;
        border-color: #60A5FA;
        color: #FFFFFF;
    }

    [data-testid="stFileUploader"] {
        background-color: #131E32;
        border: 1px solid #334155;
        border-radius: 10px;
        padding: 10px;
    }

    [data-testid="stFileUploader"] section {
        background-color: #131E32;
        border: none;
    }

    [data-testid="stFileUploader"] button {
        background-color: #2563EB;
        color: #FFFFFF;
        border: 1px solid #3B82F6;
    }

    [data-testid="stFileUploader"] small,
    [data-testid="stFileUploader"] span {
        color: #E2E8F0 !important;
    }

    [data-testid="stExpander"] {
        background-color: #131E32 !important;
        border: 1px solid #334155 !important;
        border-radius: 10px;
    }

    [data-testid="stExpander"] summary {
        background-color: #17243B !important;
        color: #FFFFFF !important;
    }

    [data-testid="stExpander"] pre,
    [data-testid="stExpander"] code {
        background-color: #0B1220 !important;
        color: #E2E8F0 !important;
    }

    .stApp label {
        color: #E2E8F0;
    }

    [data-testid="stAlert"] {
        border-radius: 10px;
    }

    .stCaption,
    [data-testid="stCaptionContainer"],
    [data-testid="stCaptionContainer"] p {
        color: #CBD5E1 !important;
    }
    </style>
    """,
    unsafe_allow_html=True,
)


# =========================================================
# HELPER FUNCTIONS
# =========================================================

def get_pdf_info(uploaded_file):
    """Return the PDF page count and size in KB."""
    pdf_bytes = uploaded_file.getvalue()

    with fitz.open(stream=pdf_bytes, filetype="pdf") as pdf:
        page_count = len(pdf)

    return page_count, len(pdf_bytes) / 1024


def file_signature(uploaded_file):
    """Detect changes to uploaded documents."""
    if uploaded_file is None:
        return None

    file_hash = hashlib.sha256(
        uploaded_file.getvalue()
    ).hexdigest()

    return (
        uploaded_file.name,
        uploaded_file.size,
        file_hash,
    )


def count_findings(result_text, section_name):
    """Count bullet points under a named Markdown heading."""
    if not result_text:
        return 0

    current_section = ""
    count = 0

    for line in result_text.splitlines():
        heading = re.match(
            r"^\s*#{1,6}\s+(.+?)\s*$",
            line,
        )

        if heading:
            current_section = heading.group(1).strip().lower()
            continue

        if current_section == section_name.lower():
            if re.match(r"^\s*[-*•]\s+\S+", line):
                count += 1

    return count


def get_section_color(section_name):
    """Choose a color based on the comparison section."""
    section_name = section_name.lower()

    if "added" in section_name:
        return "#22C55E"

    if "removed" in section_name:
        return "#EF4444"

    if "modified" in section_name:
        return "#F59E0B"

    if "similarit" in section_name:
        return "#60A5FA"

    return "#CBD5E1"


def render_comparison_results(result_text):
    """Display comparison results with section-based colors."""
    current_section = ""

    for line in result_text.splitlines():
        stripped = line.strip()

        if not stripped:
            st.write("")
            continue

        heading = re.match(
            r"^(#{1,6})\s+(.+?)\s*$",
            stripped,
        )

        if heading:
            heading_level = len(heading.group(1))
            heading_text = heading.group(2).strip()
            current_section = heading_text.lower()

            if heading_level == 1:
                st.markdown(f"## {heading_text}")
            elif heading_level == 2:
                st.markdown(f"### {heading_text}")
            else:
                st.markdown(f"#### {heading_text}")

            continue

        bullet = re.match(r"^[-*•]\s+(.+)$", stripped)

        if bullet:
            bullet_text = bullet.group(1)
            color = get_section_color(current_section)

            safe_text = html.escape(bullet_text)

            st.markdown(
                f"""
                <div class="finding"
                     style="border-left-color:{color}; color:{color};">
                    <span style="font-weight:600;">•</span>
                    {safe_text}
                </div>
                """,
                unsafe_allow_html=True,
            )

        else:
            # Escape HTML so AI-generated text is displayed safely.
            st.markdown(html.escape(stripped))


def add_formatted_text(paragraph, text):
    """Format simple bold Markdown in a Word paragraph."""
    parts = re.split(r"(\*\*.*?\*\*)", text)

    for part in parts:
        if not part:
            continue

        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        else:
            paragraph.add_run(part)


def add_page_number(paragraph):
    """Insert a Word page number field."""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run("Page ")
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(100, 116, 139)

    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


# =========================================================
# PROFESSIONAL WORD REPORT
# =========================================================

def create_word_report(result_text, original_name, updated_name):
    """Create a formatted Word comparison report."""
    doc = Document()
    section = doc.sections[0]

    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(55, 65, 81)
    normal.paragraph_format.space_after = Pt(6)

    for style_name, font_size in [
        ("Title", 28),
        ("Heading 1", 17),
        ("Heading 2", 13),
        ("Heading 3", 11),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(font_size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(31, 78, 121)

    header = section.header.paragraphs[0]
    header.text = "DOCUMIND AI | DOCUMENT INTELLIGENCE"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    for run in header.runs:
        run.font.name = "Arial"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(100, 116, 139)

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(45)

    run = paragraph.add_run("DocuMind AI")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(30)
    run.font.color.rgb = RGBColor(31, 78, 121)

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run("DOCUMENT COMPARISON REPORT")
    run.bold = True
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor(71, 85, 105)

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run(
        "AI-powered analysis of document similarities and changes"
    )
    run.italic = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph()
    doc.add_heading("Report Details", level=1)

    details_table = doc.add_table(rows=0, cols=2)
    details_table.style = "Light Shading Accent 1"

    details = [
        ("Original Document", original_name),
        ("Updated Document", updated_name),
        (
            "Generated On",
            datetime.now().strftime("%d %B %Y, %I:%M %p"),
        ),
    ]

    for label, value in details:
        cells = details_table.add_row().cells
        cells[0].text = label
        cells[1].text = str(value)

        for run in cells[0].paragraphs[0].runs:
            run.bold = True

    doc.add_paragraph()
    doc.add_heading("Comparison Findings", level=1)

    for raw_line in result_text.splitlines():
        line = raw_line.strip()

        if not line:
            continue

        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)

        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)

        elif line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=2)

        elif line.startswith(("- ", "* ", "• ")):
            paragraph = doc.add_paragraph(style="List Bullet")
            add_formatted_text(paragraph, line[2:].strip())

        elif re.match(r"^\d+[.)]\s+", line):
            content = re.sub(r"^\d+[.)]\s+", "", line)
            paragraph = doc.add_paragraph(style="List Number")
            add_formatted_text(paragraph, content)

        else:
            paragraph = doc.add_paragraph()
            add_formatted_text(paragraph, line)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = footer.add_run(
        "DocuMind AI | Confidential Comparison Report | "
    )
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(100, 116, 139)

    add_page_number(footer)

    output = BytesIO()
    doc.save(output)
    output.seek(0)

    return output.getvalue()


# =========================================================
# SESSION STATE
# =========================================================

defaults = {
    "comparison_result": None,
    "extracted_text_a": None,
    "extracted_text_b": None,
    "document_names": (None, None),
    "document_signatures": (None, None),
}

for key, default in defaults.items():
    if key not in st.session_state:
        st.session_state[key] = default


# =========================================================
# SIDEBAR
# =========================================================

with st.sidebar:
    st.markdown("# 📘 DocuMind AI")
    st.caption("Intelligent Document Analysis")

    st.markdown("---")
    st.markdown("### Workflow")
    st.markdown("1. Upload two PDF documents")
    st.markdown("2. Extract document text")
    st.markdown("3. Compare using AI")
    st.markdown("4. Download the report")

    st.markdown("---")
    st.markdown("### Capabilities")
    st.markdown("- PDF text extraction")
    st.markdown("- OCR for scanned PDFs")
    st.markdown("- AI-powered comparison")
    st.markdown("- Color-coded findings")
    st.markdown("- Word report generation")

    st.markdown("---")

    if st.button("Clear comparison", use_container_width=True):
        for key, default in defaults.items():
            st.session_state[key] = default
        st.rerun()


# =========================================================
# HERO SECTION
# =========================================================

st.markdown(
    """
    <div class="hero">
        <h1>DocuMind AI</h1>
        <p>
            Compare documents intelligently. Discover changes.
            Generate professional reports.
        </p>
        <p>
            Upload two PDF files to identify similarities,
            additions, removals, and modified information.
        </p>
    </div>
    """,
    unsafe_allow_html=True,
)


# =========================================================
# DOCUMENT UPLOAD
# =========================================================

st.markdown(
    '<div class="section-heading">Upload Documents</div>',
    unsafe_allow_html=True,
)

col_a, col_b = st.columns(2, gap="large")

with col_a:
    st.markdown("### 📄 Original Document")

    uploaded_a = st.file_uploader(
        "Upload the original PDF",
        type=["pdf"],
        key="original_pdf",
        help="Choose the original version of your document.",
    )

with col_b:
    st.markdown("### 📑 Updated Document")

    uploaded_b = st.file_uploader(
        "Upload the updated PDF",
        type=["pdf"],
        key="updated_pdf",
        help="Choose the newer version of your document.",
    )


# =========================================================
# RESET RESULTS WHEN FILES CHANGE
# =========================================================

current_signatures = (
    file_signature(uploaded_a),
    file_signature(uploaded_b),
)

if current_signatures != st.session_state.document_signatures:
    st.session_state.document_signatures = current_signatures
    st.session_state.comparison_result = None
    st.session_state.extracted_text_a = None
    st.session_state.extracted_text_b = None
    st.session_state.document_names = (None, None)


# =========================================================
# VALIDATE AND DISPLAY DOCUMENT INFORMATION
# =========================================================

valid_uploads = uploaded_a is not None and uploaded_b is not None

for uploaded_file in (uploaded_a, uploaded_b):
    if uploaded_file is not None:
        if uploaded_file.size > MAX_PDF_SIZE_BYTES:
            st.error(
                f"'{uploaded_file.name}' exceeds the "
                f"{MAX_PDF_SIZE_MB} MB upload limit. "
                "Please upload a smaller PDF."
            )
            valid_uploads = False

if uploaded_a is not None and uploaded_b is not None:
    if current_signatures[0] == current_signatures[1]:
        st.warning(
            "You uploaded the same PDF twice. "
            "Upload two different document versions to compare them."
        )
        valid_uploads = False

if valid_uploads:
    try:
        info_a = get_pdf_info(uploaded_a)
        info_b = get_pdf_info(uploaded_b)

        st.markdown(
            '<div class="section-heading">Document Overview</div>',
            unsafe_allow_html=True,
        )

        info_col_a, info_col_b = st.columns(2, gap="large")

        with info_col_a:
            st.markdown(
                f"""
                <div class="info-card">
                    <div class="info-label">Original document</div>
                    <div class="info-value">
                        {html.escape(uploaded_a.name)}
                    </div>
                    <p>{info_a[0]} pages · {info_a[1]:.1f} KB</p>
                </div>
                """,
                unsafe_allow_html=True,
            )

        with info_col_b:
            st.markdown(
                f"""
                <div class="info-card">
                    <div class="info-label">Updated document</div>
                    <div class="info-value">
                        {html.escape(uploaded_b.name)}
                    </div>
                    <p>{info_b[0]} pages · {info_b[1]:.1f} KB</p>
                </div>
                """,
                unsafe_allow_html=True,
            )

        st.success("Both PDF files are ready for processing.")

    except Exception as error:
        st.error(f"Unable to read PDF information: {error}")
        valid_uploads = False


# =========================================================
# EXTRACTION AND COMPARISON
# =========================================================

st.markdown(
    '<div class="section-heading">Analyze Documents</div>',
    unsafe_allow_html=True,
)

if not valid_uploads:
    st.info("Upload two different PDF documents to begin.")

if st.button(
    "🔍 Extract Text and Compare Documents",
    type="primary",
    use_container_width=True,
    disabled=not valid_uploads,
):
    try:
        with st.spinner("Extracting text from both documents..."):
            text_a = extract_text_from_pdf(uploaded_a)
            text_b = extract_text_from_pdf(uploaded_b)

        st.session_state.extracted_text_a = text_a
        st.session_state.extracted_text_b = text_b
        st.session_state.document_names = (
            uploaded_a.name,
            uploaded_b.name,
        )

        with st.spinner("AI is analyzing the document differences..."):
            result = compare_documents(text_a, text_b)

        if not result or not result.strip():
            raise ValueError(
                "The AI returned an empty result. Please try again."
            )

        st.session_state.comparison_result = result.strip()
        st.success("Document comparison completed successfully!")

    except Exception as error:
        st.session_state.comparison_result = None

        st.error(
            "The comparison could not be completed. "
            "Check your PDF files, API configuration, and internet connection."
        )
        st.caption(f"Technical details: {error}")


# =========================================================
# EXTRACTION PREVIEW
# =========================================================

if (
    st.session_state.extracted_text_a
    and st.session_state.extracted_text_b
):
    st.markdown(
        '<div class="section-heading">Extracted Text Preview</div>',
        unsafe_allow_html=True,
    )

    preview_a, preview_b = st.columns(2, gap="large")

    with preview_a:
        st.markdown("### Original Document")

        st.caption(
            f"{len(st.session_state.extracted_text_a):,} "
            "characters extracted"
        )

        with st.expander("View original text", expanded=False):
            st.code(
                st.session_state.extracted_text_a[:12000],
                language="text",
            )

    with preview_b:
        st.markdown("### Updated Document")

        st.caption(
            f"{len(st.session_state.extracted_text_b):,} "
            "characters extracted"
        )

        with st.expander("View updated text", expanded=False):
            st.code(
                st.session_state.extracted_text_b[:12000],
                language="text",
            )


# =========================================================
# COMPARISON RESULTS — FULL WIDTH
# =========================================================

result_text = st.session_state.comparison_result

if result_text:
    st.markdown("---")

    st.markdown(
        '<div class="section-heading">Comparison Overview</div>',
        unsafe_allow_html=True,
    )

    added_count = count_findings(result_text, "Added Content")
    removed_count = count_findings(result_text, "Removed Content")
    modified_count = count_findings(
        result_text,
        "Modified Information",
    )

    metric_col1, metric_col2, metric_col3, metric_col4 = st.columns(4)

    with metric_col1:
        st.metric("Documents Compared", "2")

    with metric_col2:
        st.metric("Added Findings", added_count)

    with metric_col3:
        st.metric("Removed Findings", removed_count)

    with metric_col4:
        st.metric("Modified Findings", modified_count)

    st.markdown(
        '<div class="section-heading">AI Comparison Results</div>',
        unsafe_allow_html=True,
    )

    # Display results using section-based colors.
    render_comparison_results(result_text)

    st.markdown(
        '<div class="section-heading">Download Your Report</div>',
        unsafe_allow_html=True,
    )

    original_name, updated_name = st.session_state.document_names

    try:
        word_data = create_word_report(
            result_text=result_text,
            original_name=original_name or "Original document",
            updated_name=updated_name or "Updated document",
        )

        st.download_button(
            label="📥 Download Professional Word Report",
            data=word_data,
            file_name="DocuMind_AI_Comparison_Report.docx",
            mime=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
            use_container_width=True,
            key="documind_word_report_download",
        )

    except Exception as error:
        st.error(f"Could not generate the Word report: {error}")


# =========================================================
# FOOTER
# =========================================================

st.markdown(
    """
    <div class="footer">
        <strong>DocuMind AI</strong><br>
        AI-Powered Document Comparison and Analysis<br>
        Built with Python, Streamlit, PyMuPDF, Tesseract OCR,
        and Gemini AI
    </div>
    """,
    unsafe_allow_html=True,
)